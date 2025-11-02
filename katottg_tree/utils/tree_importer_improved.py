# Покращена версія імпортера з детальним логуванням
import json
import os
import uuid

import frappe
from frappe import _
from frappe.utils import now_datetime
from openpyxl import load_workbook


def background_importer_improved(file_path, site, tracking_id):
	"""Покращена версія імпорту з детальним логуванням та обробкою помилок"""
	import_stats = {
		"total_rows": 0,
		"imported_rows": 0,
		"skipped_rows": 0,
		"errors": [],
		"start_time": now_datetime(),
	}

	try:
		frappe.init(site=site)
		frappe.connect()
		job_id = tracking_id

		# Отримуємо налаштування
		from katottg_tree.katottg_tree.doctype.katottg_settings.katottg_settings import (
			get_excel_column_mapping,
			get_import_batch_size,
		)

		try:
			column_mapping = get_excel_column_mapping()
			batch_size = get_import_batch_size()
		except Exception as e:
			frappe.logger().warning(f"Використовуємо значення за замовчуванням: {e!s}")
			column_mapping = {
				"level_1": 0,
				"level_2": 1,
				"level_3": 2,
				"level_4": 3,
				"level_5": 4,
				"category": 5,
				"title": 6,
			}
			batch_size = 500

		frappe.logger().info(f"Початок імпорту. Job ID: {job_id}, Batch: {batch_size}")
		update_progress_improved(job_id, 0, 100, "Підготовка до імпорту...")

		# Очищення старих даних
		old_count = frappe.db.count("KATOTTG")
		frappe.logger().info(f"Видалення {old_count} старих записів...")
		frappe.db.delete("KATOTTG")
		frappe.db.commit()

		update_progress_improved(job_id, 5, 100, "Читання файлу...")

		# Читання Excel
		workbook = load_workbook(filename=file_path, read_only=True)
		sheet = workbook.active
		total_rows = sheet.max_row - 4
		import_stats["total_rows"] = total_rows

		rows_iterator = sheet.iter_rows()
		for _i in range(4):
			next(rows_iterator)

		fields = ["name", "code", "title", "category", "parent_katottg", "is_group"]
		values_to_insert = []

		for idx, row in enumerate(rows_iterator):
			try:
				if idx > 0 and idx % batch_size == 0:
					percent = 5 + int((idx / total_rows) * 90)
					update_progress_improved(
						job_id,
						percent,
						100,
						f"Оброблено {idx}/{total_rows} (імпорт: {import_stats['imported_rows']}, пропущ: {import_stats['skipped_rows']})",
					)

				row_data = [cell.value for cell in row]
				if not any(row_data):
					import_stats["skipped_rows"] += 1
					continue

				structure = get_row_structure_improved(row_data, column_mapping)

				if not structure["current_code"] or not structure["title"]:
					import_stats["skipped_rows"] += 1
					continue

				values = [
					structure["current_code"],
					structure["current_code"],
					structure["title"],
					structure["category"],
					structure["parent_code"],
					0,
				]
				values_to_insert.append(values)
				import_stats["imported_rows"] += 1

				if len(values_to_insert) >= batch_size:
					frappe.db.bulk_insert("KATOTTG", fields, values_to_insert, ignore_duplicates=True)
					values_to_insert.clear()

			except Exception as e:
				error_msg = f"Помилка рядка {idx + 5}: {e!s}"
				frappe.logger().error(error_msg)
				import_stats["errors"].append(error_msg)
				import_stats["skipped_rows"] += 1

		# Фінальна вставка
		if values_to_insert:
			frappe.db.bulk_insert("KATOTTG", fields, values_to_insert, ignore_duplicates=True)

		update_progress_improved(job_id, 95, 100, "Оновлення ієрархії...")

		# Оновлення is_group
		frappe.db.sql("""
			UPDATE `tabKATOTTG`
			SET `is_group` = 1
			WHERE `name` IN (
				SELECT DISTINCT `parent_katottg`
				FROM `tabKATOTTG`
				WHERE `parent_katottg` IS NOT NULL
			)
		""")

		# Очищення кешу
		from katottg_tree.katottg_tree.doctype.katottg_settings.katottg_settings import (
			clear_katottg_cache,
		)

		clear_katottg_cache()

		frappe.db.commit()

		# Статистика
		import_stats["end_time"] = now_datetime()
		duration = import_stats["end_time"] - import_stats["start_time"]
		summary = f"Імпортовано: {import_stats['imported_rows']}, Пропущено: {import_stats['skipped_rows']}, Помилок: {len(import_stats['errors'])}, Час: {duration}"

		frappe.logger().info(summary)
		update_progress_improved(job_id, 100, 100, summary)

		_save_import_log_improved(import_stats, file_path)

	except Exception as e:
		frappe.db.rollback()
		frappe.logger().error(f"Критична помилка: {e!s}")
		frappe.log_error(frappe.get_traceback(), "KATOTTG Import Error")
		update_progress_improved(job_id, 0, 100, f"Помилка: {e!s}", is_error=True)
		raise
	finally:
		frappe.destroy()


def get_row_structure_improved(row_data, column_mapping):
	"""Отримати структуру рядка з використанням мапінгу"""
	codes = []
	for level in ["level_1", "level_2", "level_3", "level_4", "level_5"]:
		idx = column_mapping.get(level, 0)
		if idx < len(row_data) and row_data[idx]:
			codes.append(row_data[idx])

	if not codes:
		return {"level": 0, "current_code": None, "parent_code": None, "category": None, "title": None}

	current_code = codes[-1]
	parent_code = codes[-2] if len(codes) > 1 else None

	category_idx = column_mapping.get("category", 5)
	title_idx = column_mapping.get("title", 6)

	return {
		"level": len(codes),
		"current_code": current_code,
		"parent_code": parent_code,
		"category": row_data[category_idx] if category_idx < len(row_data) else None,
		"title": row_data[title_idx] if title_idx < len(row_data) else None,
	}


def update_progress_improved(job_id, current, total, message, is_error=False):
	"""Оновлює статус з логуванням"""
	frappe.logger().info(f"[{job_id}] {current}/{total}: {message}")
	progress = {"percent": min(current, total), "message": message, "is_error": is_error}
	frappe.cache.set_value(f"katottg_import_{job_id}", progress, expires_in_sec=600)


def _save_import_log_improved(stats, file_path):
	"""Зберігає лог імпорту"""
	try:
		log_data = {
			"file_path": file_path,
			"total_rows": stats["total_rows"],
			"imported_rows": stats["imported_rows"],
			"skipped_rows": stats["skipped_rows"],
			"error_count": len(stats["errors"]),
			"errors": stats["errors"][:100],  # Перші 100 помилок
			"start_time": str(stats["start_time"]),
			"end_time": str(stats.get("end_time")),
		}
		frappe.logger().info(f"Лог імпорту: {json.dumps(log_data, ensure_ascii=False)}")
	except Exception as e:
		frappe.logger().warning(f"Не вдалося зберегти лог: {e!s}")
